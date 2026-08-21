"""문서 텍스트 추출 (요구사항 5 / Phase 2).

지원: PDF, DOCX, TXT/MD (필수) + XLSX, PPTX (선택)

추출 결과는 ``Block`` 목록이며, 각 Block 은 페이지 번호와 제목(Heading) 정보를
함께 갖는다. 이 정보가 이후 Chunk metadata(페이지/Section)로 이어진다.

스캔(이미지) PDF는 초기 버전 대상이 아니다. 텍스트가 거의 추출되지 않으면
경고를 남기고 해당 문서를 건너뛴다(OCR은 별도 기능).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from app.config import SUPPORTED_EXTENSIONS
from app.logging_setup import get_logger

logger = get_logger("extract")

# 제목으로 인식할 패턴 (한국어/영문 지침문서에서 흔한 형태)
HEADING_PATTERNS = [
    re.compile(r"^#{1,6}\s+(?P<t>.+)$"),                      # markdown
    re.compile(r"^(?P<t>\d+(?:\.\d+){0,4}\.?\s+\S.{0,80})$"),  # 1. / 1.2.3 제목
    re.compile(r"^(?P<t>제\s?\d+\s?[장절조]\s*.{0,60})$"),      # 제3장 ...
    re.compile(r"^(?P<t>(?:Chapter|Section|Appendix|Annex|부록|별첨)\s+[\w.\-]+.{0,60})$", re.I),
    re.compile(r"^(?P<t>[A-Z][A-Z0-9 /\-&,()]{4,60})$"),        # ALL CAPS 제목
]

MAX_HEADING_LEN = 90


class UnsupportedDocument(ValueError):
    """지원하지 않는 형식이거나 텍스트를 추출할 수 없는 문서."""


@dataclass
class Block:
    """문서에서 추출한 최소 단위(문단/표행/제목)."""

    text: str
    page: int | None = None
    kind: str = "paragraph"  # paragraph | heading | table
    heading_path: list[str] = field(default_factory=list)


@dataclass
class ExtractedDocument:
    blocks: list[Block]
    page_count: int
    title: str


def detect_heading(line: str) -> str | None:
    """한 줄이 제목이면 제목 텍스트를, 아니면 None 을 반환한다."""
    stripped = line.strip()
    if not stripped or len(stripped) > MAX_HEADING_LEN:
        return None
    if stripped.endswith((".", "다.", "요.", "함.")) and not re.match(r"^\d", stripped):
        return None
    for pattern in HEADING_PATTERNS:
        match = pattern.match(stripped)
        if match:
            return match.group("t").strip()
    return None


def _apply_heading_path(current: list[str], heading: str) -> list[str]:
    """번호 깊이(1.2.3)를 이용해 계층 구조를 유지한다."""
    match = re.match(r"^(\d+(?:\.\d+)*)", heading)
    if match:
        depth = match.group(1).count(".") + 1
        path = current[: depth - 1]
        path.append(heading)
        return path
    # 번호가 없는 제목은 새로운 상위 제목으로 본다.
    return [heading]


def _lines_to_blocks(text: str, page: int | None, state: list[str]) -> list[Block]:
    blocks: list[Block] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            body = "\n".join(buffer).strip()
            if body:
                blocks.append(Block(text=body, page=page, heading_path=list(state)))
            buffer.clear()

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        heading = detect_heading(line)
        if heading:
            flush()
            state[:] = _apply_heading_path(state, heading)
            blocks.append(
                Block(text=heading, page=page, kind="heading", heading_path=list(state))
            )
        elif not line.strip():
            flush()
        else:
            buffer.append(line.strip())
    flush()
    return blocks


# ------------------------------------------------------------------ formats
def extract_txt(path: Path) -> ExtractedDocument:
    raw = None
    for encoding in ("utf-8", "utf-8-sig", "cp949", "euc-kr", "latin-1"):
        try:
            raw = path.read_text(encoding=encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if raw is None:
        raise UnsupportedDocument("텍스트 인코딩을 확인할 수 없습니다.")
    state: list[str] = []
    blocks = _lines_to_blocks(raw, page=1, state=state)
    return ExtractedDocument(blocks=blocks, page_count=1, title=path.stem)


def extract_pdf(path: Path) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 설치 안내용
        raise UnsupportedDocument(
            "PDF 처리를 위해 pypdf 가 필요합니다. pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    state: list[str] = []
    blocks: list[Block] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # pragma: no cover - 깨진 페이지는 건너뛴다
            logger.warning("PDF page extract failed | page=%s", page_number)
            continue
        blocks.extend(_lines_to_blocks(text, page=page_number, state=state))

    title = ""
    try:
        meta = reader.metadata
        if meta and meta.title:
            title = str(meta.title).strip()
    except Exception:  # pragma: no cover
        title = ""
    return ExtractedDocument(
        blocks=blocks, page_count=len(reader.pages), title=title or path.stem
    )


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedDocument(
            "DOCX 처리를 위해 python-docx 가 필요합니다. pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    state: list[str] = []
    blocks: list[Block] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading") or style in ("title", "subtitle"):
            level = 1
            match = re.search(r"(\d+)", style)
            if match:
                level = max(1, int(match.group(1)))
            state[:] = [*state[: level - 1], text]
            blocks.append(
                Block(text=text, page=None, kind="heading", heading_path=list(state))
            )
        else:
            detected = detect_heading(text)
            if detected:
                state[:] = _apply_heading_path(state, detected)
                blocks.append(
                    Block(text=text, page=None, kind="heading", heading_path=list(state))
                )
            else:
                blocks.append(Block(text=text, page=None, heading_path=list(state)))

    # 표는 행 단위로 이어붙여 하나의 블록으로 보존한다(요구사항 8 - 표 구조 보존).
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(
                Block(
                    text="\n".join(rows),
                    page=None,
                    kind="table",
                    heading_path=list(state),
                )
            )

    title = path.stem
    if document.paragraphs:
        first = document.paragraphs[0].text.strip()
        if first and len(first) <= MAX_HEADING_LEN:
            title = first
    return ExtractedDocument(blocks=blocks, page_count=0, title=title)


def extract_xlsx(path: Path) -> ExtractedDocument:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UnsupportedDocument(
            "XLSX 처리를 위해 openpyxl 이 필요합니다. pip install openpyxl"
        ) from exc

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        state = [f"[Sheet] {sheet.title}"]
        blocks.append(
            Block(text=sheet.title, page=sheet_index, kind="heading", heading_path=list(state))
        )
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(
                Block(
                    text="\n".join(rows),
                    page=sheet_index,
                    kind="table",
                    heading_path=list(state),
                )
            )
    workbook.close()
    return ExtractedDocument(
        blocks=blocks, page_count=len(workbook.sheetnames), title=path.stem
    )


def extract_pptx(path: Path) -> ExtractedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise UnsupportedDocument(
            "PPTX 처리를 위해 python-pptx 가 필요합니다. pip install python-pptx"
        ) from exc

    presentation = Presentation(str(path))
    blocks: list[Block] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        state = [f"[Slide {slide_number}]"]
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
        if texts:
            if len(texts[0]) <= MAX_HEADING_LEN:
                state = [texts[0]]
            blocks.append(
                Block(
                    text="\n".join(texts),
                    page=slide_number,
                    heading_path=list(state),
                )
            )
    return ExtractedDocument(
        blocks=blocks, page_count=len(presentation.slides), title=path.stem
    )


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".txt": extract_txt,
    ".md": extract_txt,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
}


def extract(path: Path) -> ExtractedDocument:
    """확장자에 맞는 추출기를 실행한다."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS or ext not in EXTRACTORS:
        raise UnsupportedDocument(f"지원하지 않는 형식입니다: {ext}")
    result = EXTRACTORS[ext](path)

    total_chars = sum(len(block.text) for block in result.blocks)
    if total_chars < 30:
        raise UnsupportedDocument(
            "텍스트를 추출하지 못했습니다. 스캔(이미지) 문서일 수 있으며 "
            "초기 버전에서는 OCR을 지원하지 않습니다."
        )
    return result
