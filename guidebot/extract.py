"""문서 텍스트 추출 (PDF / DOCX / TXT / MD / XLSX / PPTX).

추출 결과는 Block 목록이며, 각 Block은 페이지 번호와 구조 정보를 함께 갖는다.
Scanned PDF(이미지 PDF)는 초기 버전에서 지원하지 않고 명시적으로 오류를 낸다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from .logging_setup import get_logger
from .structure import split_lines_by_heading

logger = get_logger("extract")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".pptx"}

# 이보다 적게 추출되면 텍스트가 없는 문서로 판단한다(문서 전체 기준)
MIN_PDF_CHARS = 30


class ExtractionError(RuntimeError):
    """문서에서 텍스트를 추출하지 못했을 때 발생한다."""


class MissingDependencyError(ExtractionError):
    """문서 형식에 필요한 라이브러리가 설치되지 않았을 때 발생한다."""


@dataclass
class Block:
    """문서에서 추출한 최소 단위(문단, 표 행, 제목 등)."""

    text: str
    page: int = 1
    kind: str = "body"   # body | heading | table
    level: int = 0       # heading level (1 = 최상위)


@dataclass
class ExtractedDocument:
    blocks: list[Block]
    page_count: int
    title: str = ""
    warning: str = ""   # 등록은 되지만 사용자에게 알려야 하는 사항(예: 텍스트가 적음)

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)


def _require(module: str, package: str):
    try:
        return __import__(module)
    except ImportError as exc:  # pragma: no cover - 환경 의존
        raise MissingDependencyError(
            f"{package} 라이브러리가 필요합니다. `pip install {package}` 후 다시 시도하세요."
        ) from exc


def _split_paragraphs(text: str) -> list[str]:
    """빈 줄 기준으로 문단을 나누되, 짧은 줄(제목 등)은 그대로 유지한다."""
    chunks: list[str] = []
    for raw in re.split(r"\n\s*\n", text):
        cleaned = re.sub(r"[ \t]+", " ", raw).strip()
        if cleaned:
            chunks.append(cleaned)
    return chunks


# ---------------------------------------------------------------------------
# 형식별 추출기
# ---------------------------------------------------------------------------
def extract_txt(path: Path) -> ExtractedDocument:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp949", errors="replace")
    blocks: list[Block] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(
                Block(text=stripped.lstrip("#").strip(), page=1, kind="heading", level=level)
            )
        else:
            blocks.append(Block(text=stripped, page=1))
    if not blocks:
        raise ExtractionError("문서에서 텍스트를 찾지 못했습니다.")
    return ExtractedDocument(blocks=blocks, page_count=1)


def pdf_crypto_backend() -> tuple[str, str] | None:
    """pypdf가 실제로 사용 중인 암호화 백엔드를 반환한다: (이름, 버전).

    pypdf는 어떤 백엔드(cryptography > pycryptodome > fallback)를 쓸지
    프로세스가 처음 pypdf를 import할 때 딱 한 번만 정하고 그 결과를 캐시한다.
    그래서 서버 실행 도중 pip install을 해도, 프로세스를 완전히 재시작하기 전까지는
    반영되지 않는다. 이 함수는 지금 이 프로세스가 실제로 쓰고 있는 백엔드를 그대로
    보여줘서, "설치는 했는데 왜 안 되지?"를 구분할 수 있게 한다.
    """
    try:
        from pypdf._crypt_providers import crypt_provider
    except Exception:
        return None
    return crypt_provider


def pdf_crypto_backend_is_usable() -> bool:
    """AES로 암호화된 PDF를 읽을 수 있는 백엔드가 로드되어 있는지 확인한다."""
    backend = pdf_crypto_backend()
    return bool(backend) and backend[0] != "local_crypt_fallback"


def _pdf_read_error_message(exc: Exception) -> str:
    """pypdf 예외를 사용자가 조치할 수 있는 메시지로 바꾼다."""
    text = f"{exc.__class__.__name__}: {exc}".lower()
    if "cryptography" in text or "dependencyerror" in exc.__class__.__name__.lower():
        backend = pdf_crypto_backend()
        if backend and backend[0] != "local_crypt_fallback":
            # cryptography는 로드되어 있는데도 실패한 경우: 설치 문제가 아니다.
            return (
                f"이 PDF의 암호화 방식을 처리하지 못했습니다({exc.__class__.__name__}: {exc}). "
                "파일이 손상되었거나 지원하지 않는 암호화 방식일 수 있습니다."
            )
        return (
            "이 PDF는 암호화(AES)되어 있는데, 이를 해독할 라이브러리(cryptography)가 "
            "이 프로그램에서 아직 로드되지 않았습니다.\n\n"
            "1) 터미널에서 실행하세요:  pip install cryptography\n"
            "2) 이미 설치했다면 지금 실행 중인 GuideBot을 완전히 종료(터미널 창 닫기 또는 "
            "Ctrl+C)한 뒤 run.bat으로 다시 실행하세요.\n"
            "   (Streamlit 새로고침만으로는 반영되지 않습니다 - 파이썬 프로세스가 켜져 있는 "
            "동안 설치한 라이브러리는 그 프로세스에 자동으로 반영되지 않습니다.)\n\n"
            "설정 화면(Settings)에서 '문서 처리' 상태를 보면 지금 이 프로그램이 "
            "cryptography를 실제로 쓰고 있는지 확인할 수 있습니다."
        )
    return f"PDF를 읽는 중 오류가 발생했습니다({exc.__class__.__name__}). 파일이 손상되었는지 확인하세요."


def extract_pdf(path: Path) -> ExtractedDocument:
    pypdf = _require("pypdf", "pypdf")
    try:
        reader = pypdf.PdfReader(str(path))
    except Exception as exc:  # pypdf는 다양한 예외를 던진다
        raise ExtractionError(
            f"PDF 파일을 열 수 없습니다({exc.__class__.__name__}). "
            "파일이 손상되었는지 확인하세요."
        ) from exc

    # 암호가 걸린 PDF: 사내 문서는 빈 암호로 보호된 경우가 많으므로 먼저 시도한다.
    if getattr(reader, "is_encrypted", False):
        try:
            if not reader.decrypt(""):
                raise ExtractionError(
                    "암호로 보호된 PDF입니다. 암호를 해제한 사본을 등록해 주세요."
                )
        except ExtractionError:
            raise
        except Exception as exc:
            # decrypt("")는 인증만 확인하며, 실제 내용을 읽을 때(AES) cryptography가
            # 없으면 여기서도 DependencyError가 날 수 있다.
            raise ExtractionError(_pdf_read_error_message(exc)) from exc

    # reader.pages 에 처음 접근하는 순간 내부적으로 페이지 수/객체를 복호화한다.
    # cryptography 패키지가 없으면 여기서 pypdf.errors.DependencyError가 발생한다.
    try:
        page_count = len(reader.pages)
    except Exception as exc:
        raise ExtractionError(_pdf_read_error_message(exc)) from exc

    blocks: list[Block] = []
    failed_pages = 0
    for page_no in range(1, page_count + 1):
        try:
            page = reader.pages[page_no - 1]
            text = page.extract_text() or ""
        except Exception as exc:
            # 특정 페이지만 손상된 경우 문서 전체를 포기하지 않고 건너뛴다.
            logger.warning(
                "pdf page read failed: page=%d error=%s", page_no, exc.__class__.__name__
            )
            text = ""
            failed_pages += 1
        for paragraph in _split_paragraphs(text):
            # PDF에는 서식 정보가 없으므로 줄 단위로 제목을 추정한다.
            for piece, is_heading, level in split_lines_by_heading(paragraph):
                blocks.append(
                    Block(
                        text=piece,
                        page=page_no,
                        kind="heading" if is_heading else "body",
                        level=level,
                    )
                )

    total_chars = sum(len(b.text) for b in blocks)

    # 추출된 글자가 사실상 없을 때만 거부한다(이미지로만 이루어진 문서).
    if total_chars < MIN_PDF_CHARS:
        raise ExtractionError(
            f"이 PDF에서는 텍스트를 찾지 못했습니다(추출 {total_chars}자 / {page_count}쪽). "
            "이미지·스캔으로만 이루어진 문서로 보입니다. "
            "원본 파일(DOCX/PPTX)이 있으면 그것을 등록하거나, "
            "PDF에서 텍스트 인식(OCR)을 적용한 사본을 등록해 주세요."
        )

    # 글자 수가 적더라도 등록은 진행하되 사용자에게 알린다(그림 위주 안내문 등).
    warning = ""
    if page_count and total_chars < 20 * page_count:
        warning = (
            f"텍스트가 적게 추출되었습니다({total_chars}자 / {page_count}쪽). "
            "그림 위주 문서일 수 있어 검색 품질이 낮을 수 있습니다."
        )
    if failed_pages:
        warning = (warning + " " if warning else "") + f"읽지 못한 페이지 {failed_pages}쪽 있음."

    title = ""
    try:
        meta = reader.metadata
        if meta and meta.title:
            title = str(meta.title).strip()
    except Exception:
        title = ""
    return ExtractedDocument(
        blocks=blocks, page_count=page_count, title=title, warning=warning
    )


def extract_docx(path: Path) -> ExtractedDocument:
    docx = _require("docx", "python-docx")
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"DOCX를 열 수 없습니다: {exc.__class__.__name__}") from exc

    blocks: list[Block] = []
    page = 1
    for paragraph in document.paragraphs:
        # 명시적 페이지 나눔이 있으면 페이지 번호를 증가시킨다.
        xml = paragraph._element.xml if hasattr(paragraph, "_element") else ""
        if 'w:br' in xml and 'type="page"' in xml:
            page += 1
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith(("heading", "제목")):
            digits = re.findall(r"\d+", style)
            level = int(digits[0]) if digits else 1
            blocks.append(Block(text=text, page=page, kind="heading", level=level))
        elif style.lower() == "title":
            blocks.append(Block(text=text, page=page, kind="heading", level=1))
        else:
            blocks.append(Block(text=text, page=page))

    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(Block(text="\n".join(rows), page=page, kind="table"))

    if not blocks:
        raise ExtractionError("문서에서 텍스트를 찾지 못했습니다.")
    return ExtractedDocument(blocks=blocks, page_count=page)


def extract_xlsx(path: Path) -> ExtractedDocument:
    openpyxl = _require("openpyxl", "openpyxl")
    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        raise ExtractionError(f"XLSX를 열 수 없습니다: {exc.__class__.__name__}") from exc

    blocks: list[Block] = []
    sheet_count = len(workbook.worksheets)
    for sheet_no, sheet in enumerate(workbook.worksheets, start=1):
        blocks.append(Block(text=str(sheet.title), page=sheet_no, kind="heading", level=1))
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
            if len(rows) >= 400:  # 지나치게 큰 시트는 앞부분만 사용한다
                break
        if rows:
            blocks.append(Block(text="\n".join(rows), page=sheet_no, kind="table"))
    workbook.close()
    if not blocks:
        raise ExtractionError("문서에서 텍스트를 찾지 못했습니다.")
    return ExtractedDocument(blocks=blocks, page_count=sheet_count or 1)


def extract_pptx(path: Path) -> ExtractedDocument:
    _require("pptx", "python-pptx")
    from pptx import Presentation  # noqa: WPS433 (선택적 의존성)

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError(f"PPTX를 열 수 없습니다: {exc.__class__.__name__}") from exc

    blocks: list[Block] = []
    slide_count = 0
    for slide_no, slide in enumerate(presentation.slides, start=1):
        slide_count = slide_no
        title_text = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title_text = slide.shapes.title.text.strip()
        if title_text:
            blocks.append(Block(text=title_text, page=slide_no, kind="heading", level=1))
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text and text != title_text:
                blocks.append(Block(text=text, page=slide_no))
    if not blocks:
        raise ExtractionError("문서에서 텍스트를 찾지 못했습니다.")
    return ExtractedDocument(blocks=blocks, page_count=slide_count or 1)


_EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".txt": extract_txt,
    ".md": extract_txt,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
}


def extract_document(path: Path | str) -> ExtractedDocument:
    """확장자에 맞는 추출기를 선택해 문서를 읽는다."""
    path = Path(path)
    if not path.exists():
        raise ExtractionError(f"파일을 찾을 수 없습니다: {path.name}")
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise ExtractionError(
            f"지원하지 않는 형식입니다: {path.suffix} "
            f"(지원: {', '.join(sorted(SUPPORTED_EXTENSIONS))})"
        )
    document = extractor(path)
    if not document.title:
        document.title = path.stem
    return document
